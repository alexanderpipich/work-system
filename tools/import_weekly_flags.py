"""Недельные списки аванса из ведомости — П4.4, блок D.

Лист «Ведомость»: сотрудники со словом «еженедельно» в колонке «Номер карты» —
на еженедельном авансе. Для них на неделю --week создаются отметки
AdvanceSelection (как галочка ТМ на /admin/payouts/advances), selected_by —
суперадмин. Дальше отметки переносит carry-over П3 при первом открытии
следующей недели.

Сотрудник ищется по нормализованному ФИО (employee_name_norm); не найденные —
в отчёт, отметка всё равно ставится по ФИО (как ручная галочка). Уже
отмеченные пропускаются.

Сухой прогон по умолчанию. С --apply — запись и аудит advance.selection.import.

    python tools/import_weekly_flags.py --file "9-2 ЛОф.xlsx" --week 2026-09-28
    python tools/import_weekly_flags.py --file ... --week 2026-09-28 --apply
"""

import argparse
import os
import sys
from datetime import date, timedelta

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl                                                           # noqa: E402

from audit_helpers import create_audit_log                                # noqa: E402
from database import SessionLocal                                         # noqa: E402
from models import AdvanceSelection, AuditLog, User                       # noqa: E402
from payouts_helpers import CARRY_OVER_ACTION, week_selection             # noqa: E402
from utils import find_user_by_name, normalize_person_name, normalize_text  # noqa: E402

SHEET = "Ведомость"
NAME_HEADER = "ФИО сотрудника Исполнителя"
CARD_HEADER = "Номер карты"
WEEKLY = "еженедельно"
IMPORT_ACTION = "advance.selection.import"


def _header_index(header, title):
    for index, value in enumerate(header):
        if normalize_person_name(value) == normalize_person_name(title):
            return index
    raise ValueError(f"На листе «{SHEET}» нет колонки «{title}»")


def read_weekly_names(path):
    """ФИО со словом «еженедельно» в колонке «Номер карты», без повторов, в порядке файла."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET not in workbook.sheetnames:
        raise ValueError(f"В файле нет листа «{SHEET}»")
    rows = workbook[SHEET].iter_rows(values_only=True)
    header = next(rows)
    name_col, card_col = _header_index(header, NAME_HEADER), _header_index(header, CARD_HEADER)
    names, seen = [], set()
    for row in rows:
        name = normalize_text(row[name_col])
        if not name or normalize_person_name(row[card_col]) != WEEKLY:
            continue
        key = normalize_person_name(name)
        if key not in seen:
            seen.add(key)
            names.append(name)
    return names


def opened_later_weeks(session, week_start):
    """Следующие недели, которые уже открывались: на них carry-over больше не сработает."""
    markers = session.query(AuditLog.entity_id).filter(AuditLog.action == CARRY_OVER_ACTION).all()
    weeks = sorted({date.fromisoformat(marker) for (marker,) in markers if marker})
    return [week for week in weeks if week > week_start]


def import_weekly_flags(session, names, week_start, apply=False):
    """{added: [ФИО], existing: [ФИО], not_found: [ФИО]}. Коммит — на вызывающем."""
    admin = (session.query(User).filter(User.role == "superadmin").order_by(User.id).first()
             or session.query(User).filter(User.is_admin == True).order_by(User.id).first())  # noqa: E712
    existing = week_selection(session, week_start)
    result = {"added": [], "existing": [], "not_found": []}
    for name in names:
        user = find_user_by_name(session, name)
        if user is None:
            result["not_found"].append(name)
        if normalize_person_name(name) in existing:
            result["existing"].append(name)
            continue
        employee_name = user.employee_name if user else name
        result["added"].append(employee_name)
        if apply:
            session.add(AdvanceSelection(week_start=week_start, employee_name=employee_name,
                                         user_id=user.id if user else None,
                                         selected_by=admin.id if admin else None))
    if apply and result["added"]:
        create_audit_log(session, None, None, IMPORT_ACTION, "advance_week", week_start.isoformat(),
                         f"Аванс {week_start:%d.%m.%Y}",
                         new_value={"added": len(result["added"]), "not_found": len(result["not_found"])},
                         comment="tools/import_weekly_flags.py")
    return result


def main():
    parser = argparse.ArgumentParser(description="Отметки еженедельного аванса из ведомости")
    parser.add_argument("--file", required=True)
    parser.add_argument("--week", required=True, help="понедельник недели, ГГГГ-ММ-ДД")
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    week_start = date.fromisoformat(args.week)
    if week_start.weekday() != 0:
        monday = week_start - timedelta(days=week_start.weekday())
        print(f"{week_start:%d.%m.%Y} — не понедельник; неделя начинается {monday:%d.%m.%Y}.")
        return 1
    names = read_weekly_names(args.file)
    session = SessionLocal()
    try:
        result = import_weekly_flags(session, names, week_start, apply=args.apply)
        if args.apply:
            session.commit()
        print(f"«{WEEKLY}» в ведомости: {len(names)}")
        print(f"{'Отмечено' if args.apply else 'Будет отмечено'} на неделю {week_start:%d.%m.%Y}: "
              f"{len(result['added'])}; уже отмечены: {len(result['existing'])}")
        if result["not_found"]:
            print(f"Не найдены среди сотрудников ({len(result['not_found'])}) — отметка по ФИО из ведомости:")
            for name in result["not_found"]:
                print(f"  {name}")
        later = opened_later_weeks(session, week_start)
        if later:
            print("⚠ Уже открывались недели " + ", ".join(f"{w:%d.%m}" for w in later)
                  + " — на них отметки сами не перенесутся. Импортируйте и в текущую неделю (--week).")
        if not args.apply:
            print("Сухой прогон — --apply для записи.")
    finally:
        session.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
