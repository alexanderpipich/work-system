"""П4: документы, договоры, сроки, увольнения."""

import io
import os
import unittest
from datetime import date, timedelta
from pathlib import Path

os.environ.setdefault("DATABASE_URL", "sqlite:///:memory:")
os.environ.setdefault("SECRET_KEY", "test-secret")

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from contract_data import contract_fields, needs_patent, petition_fields, save_fixes
from contract_generator import (CONTRACT_SUBS, NO_DATA, PETITION_SUBS, Sub, TemplateMismatch, generate_contract,
                                generate_petition, render, rubles_phrase)
from database import Base
from document_alerts import reached_level, run_document_expiry_alerts
from document_helpers import alert_levels, classify_document, preview_kind
from employee_profile import employee_profile, guess_gender, parse_patent_number, patent_expiry
from models import (AuditLog, CitizenshipRegime, Country, DocumentExpiryAlert, DocumentType, EmployeeDocument,
                    LegalEntity, Region, RegimeDocumentRule, Shift, User)
from region_helpers import region_for_employee
from time_helpers import business_today

FIXTURE = Path(__file__).parent / "fixtures" / "mini_template.docx"
EMP = "Отабоева Навбахор Исломбой Кизи"


def _session():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine)


def _yellow(content):
    doc = Document(io.BytesIO(content))
    runs = [r for p in doc.paragraphs for r in p.runs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                runs += [r for p in cell.paragraphs for r in p.runs]
    return [r.text for r in runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]


def seed(session):
    patent_regime = CitizenshipRegime(code="patent", name="Патент")
    rf_regime = CitizenshipRegime(code="rf", name="РФ")
    session.add_all([patent_regime, rf_regime])
    session.flush()
    uz = Country(name="Узбекистан", regime_id=patent_regime.id, name_genitive="Республики Узбекистан")
    ru = Country(name="Россия", regime_id=rf_regime.id, name_genitive="Российской Федерации")
    session.add_all([uz, ru])
    types = {name: DocumentType(name=name, sort_order=i, requires_expiry_date=name in {"ПАТЕНТ", "РЕГИСТРАЦИЯ"},
                                alert_days="30,14,7,0" if name == "ПАТЕНТ" else None, is_active=True)
             for i, name in enumerate(["Паспорт", "СНИЛС", "ПАТЕНТ", "РЕГИСТРАЦИЯ", "ТРУДОВОЙ ДОГОВОР", "ХОДАТАЙСТВО"], 1)}
    session.add_all(types.values())
    session.flush()
    session.add(RegimeDocumentRule(regime_id=patent_regime.id, document_type_id=types["ПАТЕНТ"].id))
    session.add(LegalEntity(name="ООО «Новая»", short_name="ООО «Новая»", is_active=True,
                            full_name="Общество с ограниченной ответственностью «Новая»", inn="7805595716",
                            kpp="780601001", ogrn="1127847392243", legal_address="Москва, ул. Тестовая, 1",
                            actual_address="СПб, Лиговский 30", work_address="СПб, ул. Бабушкина 40",
                            bank_name="Банк", bank_account="40702810370110000801", bik="044525092",
                            corr_account="30101810645250000092", director_name="Иванов Иван Иванович",
                            director_name_genitive="Иванова Ивана Ивановича", director_short="И.И. Иванов",
                            director_title="генерального директора"))
    session.add(Region(name="СПб и ЛО", cities="Санкт-Петербург, Ленинградская область",
                       uvm_addressee="Начальнику УВМ ГУ МВД России по Санкт-Петербургу и Ленинградской области",
                       uvm_rank="Полковнику полиции", uvm_person=None,
                       patent_authority="ГУ МВД России по г. Санкт-Петербургу и Ленинградской области",
                       patent_territory="г. Санкт-Петербург", is_active=True))
    tm = User(id=10, employee_name="Пиунов Андрей", role="tm", phone="70000000010", password_hash="x")
    clerk = User(id=11, employee_name="Кадровик", role="hr_clerk", phone="70000000011", password_hash="x",
                 email="clerk@example.com")
    employee = User(id=1, employee_name=EMP, role="employee", tm_id=10, citizenship_country_id=None,
                    birth_date=date(1998, 3, 21), inn="780114424010", legal_entity="ООО «Новая»")
    session.add_all([tm, clerk, employee])
    session.flush()
    employee.citizenship_country_id = uz.id
    session.add(EmployeeDocument(user_id=1, employee_name=EMP, document_type_id=types["Паспорт"].id,
                                 document_number="FA9061607", issue_date=date(2024, 1, 25),
                                 extra_field_1_value="МВД РУз", status="verified"))
    session.add(EmployeeDocument(user_id=1, employee_name=EMP, document_type_id=types["ПАТЕНТ"].id,
                                 document_number="78 №2600131320 / ПР 8408114", issue_date=date(2026, 8, 18),
                                 extra_field_1_value="ГУ МВД России по г. Санкт-Петербургу и Ленинградской области",
                                 extra_field_2_value="подсобный рабочий", status="verified"))
    session.add(Shift(employee=EMP, store="ТК-1", format="ГМ", city="Санкт-Петербург",
                      shift_date=date(2026, 10, 5), service="Выкладка", hours=11, request_type="Основные заказы"))
    session.commit()
    return types


class ProfileTests(unittest.TestCase):
    def test_helpers(self):
        self.assertEqual(guess_gender(EMP), "f")
        self.assertEqual(guess_gender("Каримов Алишер Угли"), "m")
        self.assertEqual(guess_gender("Иванов Иван Сергеевич"), "m")
        self.assertIsNone(guess_gender("Смит Джон"))
        self.assertEqual(parse_patent_number("78 №2600131320 / ПР 8408114"),
                         {"series": "78", "number": "2600131320", "blank_series": "ПР", "blank_number": "8408114"})
        self.assertEqual(preview_kind("a/b.PDF"), "pdf")
        self.assertEqual(preview_kind("a/b.heic"), "none")
        self.assertEqual(rubles_phrase(232.18), "232,18 (двести тридцать два) рубля 18 копеек")
        self.assertEqual(rubles_phrase(1021), "1021,00 (одна тысяча двадцать один) рубль 00 копеек")

    def test_patent_expiry_year_minus_day(self):
        doc = EmployeeDocument(issue_date=date(2026, 8, 18))
        self.assertEqual(patent_expiry(doc), (date(2027, 8, 17), True))

    def test_profile_and_region(self):
        session = _session()()
        seed(session)
        profile = employee_profile(session, session.get(User, 1))
        self.assertEqual((profile["gender"], profile["gender_guessed"]), ("f", True))
        self.assertEqual(profile["citizenship_genitive"], "Республики Узбекистан")
        self.assertEqual((profile["patent_series"], profile["patent_blank_number"]), ("78", "8408114"))
        self.assertEqual(profile["patent_expiry"], date(2027, 8, 17))
        self.assertEqual(region_for_employee(session, EMP).name, "СПб и ЛО")
        self.assertTrue(needs_patent(session, profile))

    def test_alert_threshold(self):
        self.assertEqual(alert_levels(DocumentType(alert_days="45,10,0")), [45, 10, 0])
        doc = EmployeeDocument(status="verified", expiry_date=business_today() + timedelta(days=40))
        self.assertEqual(classify_document(doc), "ok")
        self.assertEqual(classify_document(doc, warning_days=45), "expiring_soon")


class GeneratorTests(unittest.TestCase):
    def test_fixture_runs_highlight_and_no_data(self):
        subs = [Sub("Договор №", "Иванов Иван", "fio"), Sub("Юрлицо:", "ООО «Образец»", "le")]
        content, missing = render(FIXTURE, subs, {"fio": "Петров Пётр"}, allow_missing=True)
        doc = Document(io.BytesIO(content))
        self.assertEqual(doc.paragraphs[0].text, "Договор № Петров Пётр, именуемый Работником.")
        self.assertEqual(doc.paragraphs[1].text, f"Юрлицо: {NO_DATA} (далее Работодатель).")
        self.assertEqual(_yellow(content), ["Петров Пётр", NO_DATA])  # прочая подсветка снята
        self.assertEqual(missing, ["le"])
        with self.assertRaises(ValueError):
            render(FIXTURE, subs, {"fio": "Петров Пётр"})
        with self.assertRaises(TemplateMismatch):
            render(FIXTURE, [Sub("Договор №", "Нет такого", "fio")], {"fio": "x"})

    def test_real_templates_all_samples_present(self):
        values = {s.key: f"<{s.key}>" for s in CONTRACT_SUBS}
        content, missing = generate_contract(values)
        self.assertEqual(missing, [])
        self.assertIn("<employee_party>", _yellow(content))
        values = {s.key: f"<{s.key}>" for s in PETITION_SUBS}
        content, _ = generate_petition(values)
        self.assertIn("<uvm_rank>", _yellow(content))


class ContractDataTests(unittest.TestCase):
    def setUp(self):
        self.session = _session()()
        self.types = seed(self.session)
        self.user = self.session.get(User, 1)
        self.entity = self.session.query(LegalEntity).one()
        self.region = self.session.query(Region).one()

    def _fields(self, **form):
        base = {"contract_date": date(2026, 10, 11), "contract_number": "26-0001", "rate": "232,18",
                "city": "г. Санкт-Петербург", "position": "", "term_until": None, "term_indefinite": False}
        base.update(form)
        profile = employee_profile(self.session, self.user)
        return contract_fields(self.session, profile, self.entity, self.region, base)

    def test_missing_snils_then_fix_saved_to_document(self):
        fields, values, errors, with_patent = self._fields(term_until=date(2027, 8, 17))
        self.assertTrue(with_patent)
        self.assertEqual([f["key"] for f in fields if f["missing"]], ["snils"])
        save_fixes(self.session, None, None, self.user, self.entity, self.region,
                   {"doc:СНИЛС:document_number": "229-798-702 37"})
        self.session.commit()
        snils = self.session.query(EmployeeDocument).filter(
            EmployeeDocument.document_type_id == self.types["СНИЛС"].id).one()
        self.assertEqual(snils.document_number, "229-798-702 37")
        fields, values, errors, _ = self._fields(term_until=date(2027, 8, 17))
        self.assertEqual([f for f in fields if f["missing"]], [])
        content, missing = generate_contract(values)
        yellow = _yellow(content)
        self.assertIn("гражданка Республики Узбекистан " + EMP, yellow)
        self.assertIn("на срок до 17 августа 2027 года.", yellow)
        self.assertIn("229-798-702 37", yellow)
        self.assertIn("подсобный рабочий", yellow)  # профессия из патента

    def test_term_cannot_exceed_patent(self):
        _, _, errors, _ = self._fields(term_until=date(2027, 12, 31))
        self.assertTrue(any("позже окончания патента" in e for e in errors))

    def test_russian_citizen_has_no_patent_line(self):
        self.user.citizenship_country_id = self.session.query(Country).filter(Country.name == "Россия").one().id
        self.session.commit()
        fields, values, errors, with_patent = self._fields(term_indefinite=True)
        self.assertFalse(with_patent)
        self.assertFalse(any(f["key"].startswith("patent_") for f in fields))
        content, _ = generate_contract(values, allow_missing=True, with_patent=False)
        text = "\n".join(p.text for p in Document(io.BytesIO(content)).paragraphs)
        self.assertIn("на неопределенный срок.", text)
        self.assertNotIn("Патент серия", text)

    def test_petition_missing_region_person_then_no_data(self):
        profile = employee_profile(self.session, self.user)
        fields, values = petition_fields(self.session, profile, self.entity, self.region, {"petition_date": date(2026, 10, 11)})
        self.assertEqual([f["key"] for f in fields if f["missing"]], ["uvm_person"])
        content, missing = generate_petition(values, allow_missing=True)
        self.assertEqual(missing, ["uvm_person"])
        texts = [p.text for p in Document(io.BytesIO(content)).paragraphs if p.text.strip()]
        self.assertIn("Начальнику УВМ ГУ МВД России", texts)
        self.assertIn("по Санкт-Петербургу и Ленинградской области", texts)
        self.assertIn(NO_DATA, texts)


class AlertsTests(unittest.TestCase):
    def test_levels_idempotent_three_recipients(self):
        self.assertEqual(reached_level([30, 14, 7, 0], 20), 30)
        self.assertEqual(reached_level([30, 14, 7, 0], 3), 7)
        self.assertEqual(reached_level([30, 14, 7, 0], -5), 0)
        self.assertIsNone(reached_level([30, 14, 7, 0], 31))

        session = _session()()
        types = seed(session)
        patent = session.query(EmployeeDocument).filter(EmployeeDocument.document_type_id == types["ПАТЕНТ"].id).one()
        patent.expiry_date = business_today() + timedelta(days=20)
        session.commit()
        first = run_document_expiry_alerts(session)
        self.assertEqual(first["alerts_created"], 3)
        self.assertEqual(first["emails"], 1)  # почта есть только у кадровика
        self.assertEqual(run_document_expiry_alerts(session)["alerts_created"], 0)
        recipients = {a.recipient_user_id for a in session.query(DocumentExpiryAlert).all()}
        self.assertEqual(recipients, {1, 10, 11})
        self.assertEqual({a.level for a in session.query(DocumentExpiryAlert).all()}, {30})
        # Новый патент снимает тревогу по старому.
        session.add(EmployeeDocument(user_id=1, employee_name=EMP, document_type_id=types["ПАТЕНТ"].id,
                                     document_number="78 №1", issue_date=business_today(),
                                     expiry_date=business_today() + timedelta(days=300), status="uploaded"))
        session.commit()
        patent.expiry_date = business_today() + timedelta(days=5)
        session.commit()
        self.assertEqual(run_document_expiry_alerts(session)["alerts_created"], 0)


class RouteTests(unittest.TestCase):
    def setUp(self):
        import main
        from dependencies import current_user, get_db

        self.Session = _session()
        session = self.Session()
        self.types = {name: doc_type.id for name, doc_type in seed(session).items()}
        session.add(User(id=20, employee_name="Админ", role="superadmin", is_admin=True, phone="70000000020",
                         password_hash="x"))
        session.commit()

        def override_db():
            db = self.Session()
            try:
                yield db
            finally:
                db.close()

        self.app, self.get_db, self.current_user = main.app, get_db, current_user
        self.app.dependency_overrides[get_db] = override_db
        self.client = TestClient(main.app)

    def tearDown(self):
        self.app.dependency_overrides.pop(self.get_db, None)
        self.app.dependency_overrides.pop(self.current_user, None)

    def _as(self, user_id):
        db = self.Session()
        user = db.get(User, user_id)
        db.expunge(user)
        self.app.dependency_overrides[self.current_user] = lambda: user

    def test_registry_filter_accepts_typed_name(self):
        self._as(20)
        html = self.client.get("/admin/documents", params={"employee_name": EMP, "status": "all"}).text
        self.assertIn('list="documents-filter-employees"', html)
        self.assertIn(f'value="{EMP}"', html)
        self.assertIn("ПАТЕНТ", html)
        other = self.client.get("/admin/documents", params={"employee_name": "Нет Такого", "status": "all"}).text
        self.assertNotIn("FA9061607", other)

    def test_bulk_verify_partial(self):
        db = self.Session()
        ids = []
        for status in ("uploaded", "archived", "pending_verification"):
            doc = EmployeeDocument(user_id=1, employee_name=EMP, document_type_id=self.types["РЕГИСТРАЦИЯ"],
                                   status=status)
            db.add(doc)
            db.flush()
            ids.append(doc.id)
        db.commit()
        self._as(20)
        response = self.client.post("/admin/documents/verify-bulk", data={"document_ids": ids + [99999]},
                                    follow_redirects=False)
        self.assertIn("message=", response.headers["location"])
        db = self.Session()
        statuses = [db.get(EmployeeDocument, i).status for i in ids]
        self.assertEqual(statuses, ["verified", "archived", "verified"])
        audits = db.query(AuditLog).filter(AuditLog.action == "documents.verify_bulk").all()
        self.assertEqual(len(audits), 1)

    def test_contract_generated_as_document(self):
        self._as(11)  # кадровик
        response = self.client.post("/admin/contracts/contract/1", data={
            "contract_number": "26-0007", "contract_date": "2026-10-11", "city": "г. Санкт-Петербург",
            "rate": "232,18", "term_indefinite": "0", "term_until": "2027-08-17", "action": "generate",
            "fix:doc:СНИЛС:document_number": "229-798-702 37"})
        self.assertEqual(response.status_code, 200)
        self.assertIn("filename*=UTF-8''", response.headers["content-disposition"])
        db = self.Session()
        contract = db.query(EmployeeDocument).filter(
            EmployeeDocument.document_type_id == self.types["ТРУДОВОЙ ДОГОВОР"]).one()
        self.assertEqual((contract.document_number, contract.expiry_date, contract.extra_field_1_value),
                         ("26-0007", date(2027, 8, 17), "ООО «Новая»"))

    def test_contracts_rights(self):
        self._as(10)  # ТМ — просмотр, не создание
        self.assertEqual(self.client.get("/admin/contracts", follow_redirects=False).status_code, 200)
        self.assertEqual(self.client.get("/admin/contracts/contract/1", follow_redirects=False).status_code, 302)
        self._as(1)  # сотрудник
        self.assertEqual(self.client.get("/admin/contracts", follow_redirects=False).status_code, 302)

    def test_dismissal_rule_contract_without_shifts(self):
        db = self.Session()
        db.add(User(id=2, employee_name="Без Смен Пётр", role="employee"))
        db.add(EmployeeDocument(user_id=2, employee_name="Без Смен Пётр", document_number="26-0009",
                                document_type_id=self.types["ТРУДОВОЙ ДОГОВОР"], issue_date=date(2026, 9, 1),
                                is_permanent=True, status="uploaded"))
        db.commit()
        self._as(20)
        html = self.client.get("/admin/employees/dismissals").text
        self.assertIn("Без Смен Пётр", html)
        self.assertIn("договор №26-0009", html)
        self.assertIn("смен нет", html)


if __name__ == "__main__":
    unittest.main()
