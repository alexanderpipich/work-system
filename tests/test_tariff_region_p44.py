"""П4.4: тарифный регион магазина, надбавка к часу, кабинеты ТМ, недельные списки,
выплатной табель в формате ведомости."""

import datetime
import os
import unittest

os.environ.setdefault("DATABASE_URL", "sqlite:///:memory:")
os.environ.setdefault("SECRET_KEY", "test-secret")

from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from database import Base
from models import (AdvanceSelection, AuditLog, HourlyBonus, PayrollRun, PayrollRunItem, Rate, Service, Shift,
                    Store, User)
from payroll_run_helpers import build_run_items
from service_catalog import build_service_resolver
from store_helpers import populate_stores_from_shifts
from utils import pick_rate

DAY = datetime.date(2026, 9, 16)


def _session():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine)()


def _shift(store="Лента-660 (г. Одинцово)", city="Москва", service="3ур_Универсальная услуга для ММФ",
           hours=11.0, employee="Абазов Жанторо", fmt="СМ", day=DAY, request_type="Основные заказы"):
    return Shift(employee=employee, store=store, city=city, service=service, format=fmt, hours=hours,
                 shift_date=day, request_type=request_type)


def _grid(session):
    service = Service(name="Универсальная услуга для ММФ", is_active=True)
    session.add(service)
    session.flush()
    for region, rate in (("Москва", 304), ("Московская область", 311)):
        session.add(Rate(service="3ур_Универсальная услуга для ММФ", service_id=service.id, level=3,
                         format="СМ", city=region, hourly_rate=rate))
    session.commit()


class RateRegionTests(unittest.TestCase):
    def setUp(self):
        self.s = _session()
        _grid(self.s)

    def _rate(self, shift):
        return pick_rate(self.s.query(Rate).all(), shift, resolve=build_service_resolver(self.s)).hourly_rate

    def test_store_region_beats_shift_city(self):
        self.s.add(Store(tk_number=660, city="Москва", rate_region="Московская область", format="СМ"))
        self.s.commit()
        self.assertEqual(self._rate(_shift()), 311)  # подмосковный ТК — ставка МО

    def test_without_region_city_of_shift(self):
        self.s.add(Store(tk_number=660, city="Москва", rate_region=None, format="СМ"))
        self.s.commit()
        self.assertEqual(self._rate(_shift()), 304)
        self.assertEqual(self._rate(_shift(store="Лента-9999")), 304)  # ТК нет в справочнике

    def test_without_resolver_old_behaviour(self):
        self.s.add(Store(tk_number=660, rate_region="Московская область"))
        self.s.commit()
        self.assertEqual(pick_rate(self.s.query(Rate).all(), _shift()).hourly_rate, 304)

    def test_populate_sets_region_to_city(self):
        self.s.add(_shift(store="Лента-777", city="Санкт-Петербург"))
        self.s.commit()
        populate_stores_from_shifts(self.s)
        self.assertEqual(self.s.query(Store).filter_by(tk_number=777).one().rate_region, "Санкт-Петербург")


class TariffPreviewTests(unittest.TestCase):
    def test_new_services_and_catalog_spelling(self):
        from routers.rates import _services_vs_catalog

        s = _session()
        s.add_all([Service(name="Табакошоп Универсальные услуги"), Service(name="Услуги кассового сервиса")])
        s.commit()
        rows = [{"service_name": "Табакошоп_Универсальные услуги"}, {"service_name": "Услуги кассового сервиса"},
                {"service_name": "Услуги по выкладке СОФ"}]
        new, spelling = _services_vs_catalog(s, rows)
        self.assertEqual(new, ["Табакошоп_Универсальные услуги", "Услуги по выкладке СОФ"])
        self.assertEqual(spelling, [("Табакошоп_Универсальные услуги", "Табакошоп Универсальные услуги")])


class TmCabinetTests(unittest.TestCase):
    def test_grant_service_logins(self):
        from tools.grant_tm_cabinets import grant_tm_cabinets
        from utils import verify_password

        s = _session()
        s.add_all([
            User(id=1, employee_name="Пиунов Андрей", role="tm"),
            User(id=2, employee_name="Амин", role="hr_manager"),
            User(id=3, employee_name="Донёр", role="tm", phone="79001112233", password_hash="x"),
            User(id=4, employee_name="Сотрудник", role="employee"),
            User(id=5, employee_name="Занял номер", role="employee", phone="70000000001"),
        ])
        s.commit()
        self.assertEqual([(u.id, p, pw) for u, p, pw in grant_tm_cabinets(s)],
                         [(1, "70000000002", None), (2, "70000000003", None)])
        self.assertIsNone(s.get(User, 1).phone)  # сухой прогон
        result = grant_tm_cabinets(s, apply=True)
        s.commit()
        amin = s.get(User, 2)
        self.assertEqual((amin.employee_name, amin.phone, amin.password_is_temporary), ("Амин", "70000000003", True))
        password = result[1][2]
        self.assertEqual(len(password), 8)
        self.assertTrue(verify_password(password, amin.password_hash))
        logs = s.query(AuditLog).filter_by(action="user.cabinet_granted").all()
        self.assertEqual(len(logs), 2)
        self.assertFalse(any(password in (log.new_value or "") for log in logs))
        self.assertEqual(grant_tm_cabinets(s), [])


class WeeklyFlagsTests(unittest.TestCase):
    def test_import_from_statement(self):
        import tempfile

        import openpyxl

        from tools.import_weekly_flags import import_weekly_flags, read_weekly_names

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Ведомость"
        sheet.append(["ФИО сотрудника Исполнителя", "ИЧ", "Номер карты"])
        sheet.append(["Абазов Жанторо", 11, "еженедельно"])
        sheet.append(["Иванов  Иван", 8, "Еженедельно "])
        sheet.append(["Петров Пётр", 8, "2200123412341234"])
        sheet.append(["Итог", 27, None])
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "v.xlsx")
            workbook.save(path)
            names = read_weekly_names(path)
        self.assertEqual(names, ["Абазов Жанторо", "Иванов  Иван"])

        s = _session()
        s.add_all([User(id=1, employee_name="Админ", role="superadmin"), User(id=2, employee_name="Иванов Иван")])
        s.commit()
        week = datetime.date(2026, 9, 28)
        result = import_weekly_flags(s, names, week, apply=True)
        s.commit()
        self.assertEqual((result["added"], result["not_found"]), (["Абазов Жанторо", "Иванов Иван"], ["Абазов Жанторо"]))
        rows = {r.employee_name: (r.user_id, r.selected_by) for r in s.query(AdvanceSelection).filter_by(week_start=week)}
        self.assertEqual(rows, {"Абазов Жанторо": (None, 1), "Иванов Иван": (2, 1)})
        self.assertEqual(s.query(AuditLog).filter_by(action="advance.selection.import").count(), 1)
        again = import_weekly_flags(s, names, week, apply=True)
        self.assertEqual((again["added"], len(again["existing"])), ([], 2))


class HourlyBonusTests(unittest.TestCase):
    def setUp(self):
        self.s = _session()
        _grid(self.s)
        self.s.add(Store(tk_number=522, city="Москва", rate_region="Москва", format="СМ"))
        self.s.add(HourlyBonus(name="Акция", amount_per_hour=5.0, active_from=datetime.date(2026, 9, 1)))
        self.shifts = [_shift(store="Лента-522"),
                       _shift(store="Лента-522", day=DAY + datetime.timedelta(days=1),
                              request_type="Смена без плана"),
                       _shift(store="Лента-522", service="3ур_Нет такой", day=DAY + datetime.timedelta(days=2))]
        self.s.add_all(self.shifts)
        self.s.commit()

    def _run(self, kind, link_paid):
        run = PayrollRun(status="draft", kind=kind, date_from=DAY, date_to=DAY + datetime.timedelta(days=6),
                         legal_entity="ООО")
        self.s.add(run)
        self.s.flush()
        items = build_run_items(self.s, run, self.shifts, date_from=run.date_from, date_to=run.date_to,
                                rates=self.s.query(Rate).all(), resolve=build_service_resolver(self.s),
                                apply_lmk=False, apply_auto=False, manual=None, link_paid=link_paid)
        run.status = "fixed"
        self.s.commit()
        return items

    def test_abazov_advance_then_period(self):
        advance = self._run("advance", link_paid=False)
        paid = advance[0]
        self.assertEqual((paid.amount, paid.bonus_amount, paid.total_amount), (3344, 55, 3399))
        self.assertEqual([i.bonus_amount for i in advance[1:]], [0, 0])  # без плана, без ЧТС
        period = self._run("period", link_paid=True)
        self.assertEqual((period[0].advanced_amount, period[0].total_amount), (3399, 0))  # без ложного остатка

    def test_closed_bonus_not_applied(self):
        bonus = self.s.query(HourlyBonus).one()
        bonus.active_to = DAY - datetime.timedelta(days=1)
        self.s.commit()
        item = self._run("advance", link_paid=False)[0]
        self.assertEqual((item.bonus_amount, item.total_amount), (0, 3344))

    def test_bonuses_sum(self):
        self.s.add(HourlyBonus(name="Ещё", amount_per_hour=2.5, active_from=DAY))
        self.s.commit()
        self.assertEqual(self._run("advance", link_paid=False)[0].bonus_amount, 11 * 7.5)

    def test_screen_adds_and_closes(self):
        import main
        from fastapi.testclient import TestClient

        from dependencies import current_user, get_db

        admin = User(id=1, employee_name="Админ", role="superadmin", is_admin=True)
        self.s.add(admin)
        self.s.commit()
        self.s.refresh(admin)
        self.s.expunge(admin)
        main.app.dependency_overrides[get_db] = lambda: self.s
        main.app.dependency_overrides[current_user] = lambda: admin
        try:
            client = TestClient(main.app)
            response = client.post("/admin/payouts/bonuses/add", follow_redirects=False,
                                   data={"name": "Ночь", "amount_per_hour": "10,5", "active_from": "2026-10-01"})
            self.assertIn("message=", response.headers["location"])
            bonus = self.s.query(HourlyBonus).filter_by(name="Ночь").one()
            self.assertEqual(bonus.amount_per_hour, 10.5)
            client.post(f"/admin/payouts/bonuses/{bonus.id}/close", data={"active_to": "2026-10-31"})
            self.assertEqual(bonus.active_to, datetime.date(2026, 10, 31))
            html = client.get("/admin/payouts/bonuses").text
            self.assertIn("Ночь", html)
            self.assertIn("31.10.2026", html)
        finally:
            main.app.dependency_overrides.pop(get_db, None)
            main.app.dependency_overrides.pop(current_user, None)


class StatementExportTests(unittest.TestCase):
    def test_two_sheets_two_tabel_rows_one_statement_row(self):
        from io import BytesIO

        import openpyxl

        from statement_export import STATEMENT_COLUMNS, build_statement_xlsx, statement_filename, statement_rows

        s = _session()
        s.add_all([User(id=1, employee_name="Абрамов Александр", role="tm"),
                   User(id=2, employee_name="Абазов Жанторо", tm_id=1),
                   User(id=3, employee_name="Борисова Вера"),
                   Store(tk_number=660, city="Москва", rate_region="Московская область", format="СМ"),
                   Shift(employee="Абазов Жанторо", store="Лента-660 (г. Одинцово)", city="Москва", format="СМ",
                         service="3ур_ММФ", hours=11, shift_date=datetime.date(2026, 9, 16),
                         request_type="Основные заказы")])
        run = PayrollRun(kind="period", status="fixed", date_from=datetime.date(2026, 9, 16),
                         date_to=datetime.date(2026, 9, 30), legal_entity="ООО")
        s.add(run)
        s.flush()

        def item(name, service, day, hours, rate, **extra):
            amount = hours * rate
            values = dict(run_id=run.id, employee_name=name, store="Лента-660 (г. Одинцово)", service=service,
                          shift_date=datetime.date(2026, 9, day), hours=hours, rate=rate, amount=amount,
                          bonus_amount=hours * 5, advanced_amount=0, lmk_amount=0, other_amount=0,
                          auto_adjustments_amount=0, manual_adjustments_amount=0)
            values.update(extra)
            values["total_amount"] = (amount + values["bonus_amount"] + values["other_amount"]
                                      - values["lmk_amount"] - values["advanced_amount"])
            return PayrollRunItem(**values)

        s.add_all([
            item("Абазов Жанторо", "3ур_ММФ", 16, 11, 304, advanced_amount=3399),
            item("Абазов Жанторо", "3ур_ММФ", 17, 10, 304),
            item("Абазов Жанторо", "2ур_Касса", 18, 8, 280),
            item("Борисова Вера", "3ур_ММФ", 16, 8, 304, other_amount=-500, manual_adjustments_amount=-500),
        ])
        s.add(AdvanceSelection(week_start=datetime.date(2026, 9, 14), employee_name="Абазов Жанторо"))
        s.commit()

        days, tabel, statement = statement_rows(s, run)
        self.assertEqual(days, [str(d) for d in range(16, 31)])
        abazov = [r for r in tabel if r["ФИО сотрудника Исполнителя"] == "Абазов Жанторо"]
        self.assertEqual(len(abazov), 2)
        mmf = next(r for r in abazov if r["Услуга"] == "3ур_ММФ")
        self.assertEqual((mmf["Номер ТК"], mmf["Город"], mmf["Тип"], mmf["16"], mmf["17"], mmf["ИЧ"], mmf["ТМ"],
                          mmf["Ставка"], mmf["ФОТ"]),
                         (660, "Московская область", "СМ", 11, 10, 21, "Абрамов Александр", 304, 6384))
        self.assertEqual(len(statement), 2)
        row = statement[0]
        self.assertEqual((row["ФИО сотрудника Исполнителя"], row["ИЧ"], row["ФОТ"], row["Аванс"], row["Акция"],
                          row["Итого к Выплате"], row["Номер карты"], row["Комментарии"]),
                         ("Абазов Жанторо", 29, 8624, 3399, 145, 8624 + 145 - 3399, "еженедельно", "еженедельно"))
        borisova = statement[1]
        self.assertEqual((borisova["Штраф"], borisova["Премия"], borisova["Итого к Выплате"]),
                         (500, None, 8 * 304 + 40 - 500))
        self.assertNotEqual(borisova["Номер карты"], "еженедельно")

        workbook = openpyxl.load_workbook(BytesIO(build_statement_xlsx(days, tabel, statement)))
        self.assertEqual(workbook.sheetnames, ["Табель", "Ведомость"])
        header = [c.value for c in workbook["Табель"][1]]
        self.assertEqual(header[:5] + header[-4:],
                         ["ФИО сотрудника Исполнителя", "Номер ТК", "Город", "Тип", "Услуга", "ИЧ", "ТМ", "Ставка",
                          "ФОТ"])
        self.assertEqual(header[5:-4], [str(d) for d in range(16, 31)])
        self.assertEqual([c.value for c in workbook["Ведомость"][1]], STATEMENT_COLUMNS)
        self.assertEqual(workbook["Ведомость"].max_row, 4)  # заголовок + 2 + «Итог»
        self.assertEqual(statement_filename(run), "Ведомость_16.09-30.09.2026.xlsx")


if __name__ == "__main__":
    unittest.main()
