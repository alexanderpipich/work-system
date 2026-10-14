"""П4.11: правки перед демо — договор при патенте, перепроверка формы при смене юрлица,
рекомендательное письмо, логотип → рабочий стол, кнопки фиксации аванса сверху и снизу,
права HR-руководителя на табель и запрет безвозвратных удалений."""

import io
import json
import os
import re
import shutil
import subprocess
import unittest
from datetime import date

os.environ.setdefault("DATABASE_URL", "sqlite:///:memory:")
os.environ.setdefault("SECRET_KEY", "test-secret")

from docx import Document
from fastapi import Request
from fastapi.testclient import TestClient
from sqlalchemy.orm import sessionmaker

from contract_data import signer_title
from contract_generator import RECOMMENDATION_SUBS, RECOMMENDATION_TEMPLATE
from models import Country, DocumentType, EmployeeDocument, LegalEntity, PayrollRun, Rate, User
from test_documents_contracts import EMP, SHOP_ADDRESS, _session, _yellow, seed as seed_contracts
from test_payouts import WEEK, A, _engine, seed as seed_payouts

UI_JS = os.path.join(os.path.dirname(__file__), "..", "static", "ui.js")
PATENT_EXPIRY = date(2027, 8, 17)  # выдан 18.08.2026, срок не указан — год минус день


def _text(content):
    doc = Document(io.BytesIO(content))
    return "\n".join(p.text for p in doc.paragraphs)


class _Client(unittest.TestCase):
    seed = None

    def setUp(self):
        import main
        from dependencies import current_user, get_db

        self.Session = self.make_session()
        session = self.Session()
        type(self).seed(session)
        session.add(User(id=30, employee_name="Петрова HR", role="hr_lead", phone="70000000030", password_hash="x"))
        session.add(User(id=40, employee_name="Суперадмин", role="superadmin", is_admin=True, phone="70000000040",
                         password_hash="x"))
        session.commit()

        def override_db():
            db = self.Session()
            try:
                yield db
            finally:
                db.close()

        self.app, self.get_db, self.current_user = main.app, get_db, current_user
        self.app.dependency_overrides[get_db] = override_db
        self.client = TestClient(main.app)

    def tearDown(self):
        self.app.dependency_overrides.pop(self.get_db, None)
        self.app.dependency_overrides.pop(self.current_user, None)

    def _as(self, user_id):
        db = self.Session()
        user = db.get(User, user_id)
        db.expunge(user)

        def override(request: Request):
            request.state.user = user  # как настоящий current_user — для бокового меню
            return user

        self.app.dependency_overrides[self.current_user] = override


class ContractFormTests(_Client):
    seed = staticmethod(seed_contracts)

    def make_session(self):
        return _session()

    def setUp(self):
        super().setUp()
        self._as(30)

    def _russian(self):
        db = self.Session()
        db.get(User, 1).citizenship_country_id = db.query(Country).filter(Country.name == "Россия").one().id
        db.commit()

    # A1
    def test_patent_term_only_until_patent_expiry(self):
        html = self.client.get("/admin/contracts/contract/1").text
        self.assertNotIn("бессрочный", html)
        self.assertNotIn('name="term_indefinite"', html)
        self.assertIn("Срочный, до даты", html)
        self.assertIn(f'name="term_until" value="{PATENT_EXPIRY.isoformat()}"', html)
        self.assertIn(f'max="{PATENT_EXPIRY.isoformat()}"', html)

    def test_without_patent_both_variants(self):
        self._russian()
        html = self.client.get("/admin/contracts/contract/1").text
        self.assertIn("бессрочный", html)
        self.assertEqual(html.count('name="term_indefinite"'), 2)
        self.assertNotIn("Срочный, до даты", html)

    def test_indefinite_with_patent_refused(self):
        response = self.client.post("/admin/contracts/contract/1", data={
            "contract_number": "26-0007", "contract_date": "2026-10-11", "rate": "232,18",
            "term_indefinite": "1", "action": "generate_skip"})
        self.assertEqual(response.status_code, 200)
        self.assertIn("При патенте договор срочный", response.text)
        self.assertEqual(self.Session().query(EmployeeDocument).filter(
            EmployeeDocument.document_number == "26-0007").count(), 0)

    # A2
    def test_check_with_other_legal_entity_keeps_input(self):
        db = self.Session()
        other = LegalEntity(name="ООО «Другая»", short_name="ООО «Другая»", is_active=True, inn="7727404143",
                            full_name="Общество с ограниченной ответственностью «Другая»", ogrn="1197746013068")
        db.add(other)
        db.commit()
        html = self.client.post("/admin/contracts/contract/1", data={
            "legal_entity_id": str(other.id), "contract_number": "26-0042", "contract_date": "2026-10-11",
            "position": "кассир", "rate": "250", "term_until": "2027-05-01", "action": "check"}).text
        self.assertIn("Общество с ограниченной ответственностью «Другая»", html)
        self.assertIn("7727404143", html)
        self.assertNotIn("Общество с ограниченной ответственностью «Новая»", html)
        self.assertIn(f'<option value="{other.id}" selected>', html)
        for value in ('value="26-0042"', 'value="кассир"', 'value="250"', 'value="2027-05-01"'):
            self.assertIn(value, html)
        self.assertIn('data-recheck="le"', html)
        self.assertIn('data-recheck="region"', html)

    # A3
    def test_recommendation_button_and_form(self):
        self.assertRegex(self.client.get("/admin/contracts").text,
                         r'href="/admin/contracts/recommendation/1"( data-tour="contracts-recommendation")?>Создать рекомендацию')
        html = self.client.get("/admin/contracts/recommendation/1").text
        self.assertIn("Рекомендательное письмо", html)
        self.assertIn('name="recommendation_date"', html)
        self.assertIn('value="г. Санкт-Петербург"', html)  # город — из региона по последним сменам
        self.assertIn('value="подсобный рабочий"', html)  # должность — из патента
        self.assertNotIn('name="region_id"', html)
        self.assertNotIn('name="contract_number"', html)

    def test_recommendation_generated_as_document(self):
        response = self.client.post("/admin/contracts/recommendation/1", data={
            "city": "г. Санкт-Петербург", "position": "работник торгового зала",
            "recommendation_date": "2026-10-13", "action": "generate"})
        self.assertEqual(response.status_code, 200, response.text[:300])
        self.assertIn("filename*=UTF-8''", response.headers["content-disposition"])
        text = _text(response.content)
        self.assertIn("Общество с ограниченной ответственностью «Новая»", text)
        self.assertIn("ОГРН 1127847392243 КПП 780601001 ИНН 7805595716", text)
        self.assertIn(SHOP_ADDRESS, text)  # шапка — адрес магазина смен, не actual_address юрлица
        self.assertNotIn("СПб, Лиговский 30", text)
        self.assertIn(f"Настоящим подтверждаем, что сотрудник {EMP} работал в г. Санкт-Петербург "
                      f"в должности «работник торгового зала».", text)
        self.assertEqual(text.count(EMP), 1)
        self.assertIn("Генеральный директор ООО «Новая»", text)
        self.assertIn("Иванов Иван Иванович", text)
        self.assertIn("«13» октября 2026 г.", text)
        for word in ("Рустам", "Рабимов", "АЛЬЯНС", "Шестак", " Он ", " его ", " он "):
            self.assertNotIn(word, text)
        self.assertNotIn("НЕТ ДАННЫХ", _yellow(response.content))
        db = self.Session()
        doc_type = db.query(DocumentType).filter(DocumentType.name == "РЕКОМЕНДАТЕЛЬНОЕ ПИСЬМО").one()
        document = db.query(EmployeeDocument).filter(EmployeeDocument.document_type_id == doc_type.id).one()
        self.assertEqual((document.user_id, document.issue_date, document.expiry_date, document.extra_field_1_value),
                         (1, date(2026, 10, 13), None, "ООО «Новая»"))

    def test_recommendation_missing_director_flagged(self):
        db = self.Session()
        db.query(LegalEntity).one().director_name = None
        db.commit()
        html = self.client.get("/admin/contracts/recommendation/1").text
        self.assertIn("не хватает: 1", html)
        self.assertIn('name="fix:le:director_name"', html)

    def test_recommendation_template_samples_present(self):
        doc = Document(str(RECOMMENDATION_TEMPLATE))
        text = "\n".join(p.text for p in doc.paragraphs)
        for sub in RECOMMENDATION_SUBS:
            with self.subTest(key=sub.key):
                self.assertIn(sub.sample, text)
        for word in ("Рустам", "Рабимов", " Он ", " его "):
            self.assertNotIn(word, text)

    def test_signer_title(self):
        self.assertEqual(signer_title("генерального директора"), "Генеральный директор")
        self.assertEqual(signer_title("директора"), "Директор")
        self.assertEqual(signer_title(""), "Директор")

    # B1
    def test_logo_leads_to_dashboard(self):
        html = self.client.get("/admin/contracts").text
        brand = re.search(r'<a href="([^"]*)" class="rail-brand"([^>]*)>', html)
        self.assertEqual(brand.group(1), "/hr")
        self.assertNotIn("data-rail-top", brand.group(2))
        self.assertIn('<a href="/hr" class="brand-home"', html)
        dashboard = self.client.get("/hr").text
        brand = re.search(r'<a href="([^"]*)" class="rail-brand"([^>]*)>', dashboard)
        self.assertEqual(brand.group(1), "/hr")
        self.assertIn("data-rail-top", brand.group(2))  # на самом дашборде — прокрутка наверх

    # C
    def test_hr_lead_cannot_hard_delete_legal_entity(self):
        entity_id = self.Session().query(LegalEntity).one().id
        for path, data in (("/admin/legal-entities/delete", {"entity_id": entity_id, "hard_delete": "1"}),
                           ("/admin/legal-entities/remove", {"entity_id": entity_id})):
            with self.subTest(path=path):
                self.client.post(path, data=data, follow_redirects=False)
                self.assertIsNotNone(self.Session().get(LegalEntity, entity_id))
        self._as(40)  # суперадмин удаляет
        self.client.post("/admin/legal-entities/delete", data={"entity_id": entity_id, "hard_delete": "1"},
                         follow_redirects=False)
        self.assertIsNone(self.Session().get(LegalEntity, entity_id))


class AdvancesAndPayrollTests(_Client):
    seed = staticmethod(seed_payouts)

    def make_session(self):
        return sessionmaker(bind=_engine())

    # B4
    def _advances(self):
        return self.client.get(f"/admin/payouts/advances?week={WEEK.isoformat()}").text

    def test_fix_actions_top_and_bottom(self):
        self._as(30)
        html = self._advances()
        self.assertIn('data-run-actions="top"', html)
        self.assertIn('data-run-actions="bottom"', html)
        self.assertEqual(html.count('action="/admin/payouts/advances/fix"'), 2)
        self.client.post("/admin/payouts/advances/select", data={"week": WEEK.isoformat(), "employee_name": A, "checked": "1"})
        self.client.post("/admin/payouts/advances/fix", data={"week": WEEK.isoformat()})
        html = self._advances()
        self.assertEqual(html.count("/unfix"), 2)
        self.assertEqual(html.count('action="/admin/payroll/runs/send"'), 2)
        self.assertEqual(html.count('action="/admin/payouts/advances/fix"'), 0)

    def test_no_fix_actions_without_permission(self):
        self._as(10)  # ТМ: нет payouts.advances.fix
        html = self._advances()
        for marker in ('action="/admin/payouts/advances/fix"', "/unfix", "/unsend"):
            self.assertNotIn(marker, html)

    # C
    def test_hr_lead_payroll_any_period(self):
        self._as(30)
        response = self.client.post("/admin/payroll/fix", data={
            "date_from": "2026-10-01", "date_to": "2026-10-07", "legal_entity": "ООО Новая", "comment": "демо"},
            follow_redirects=False)
        self.assertEqual(response.status_code, 302)
        run = self.Session().query(PayrollRun).one()
        self.assertEqual((run.date_from, run.date_to), (date(2026, 10, 1), date(2026, 10, 7)))
        self.assertIn("[произвольный период]", run.comment)

    def test_hr_lead_cannot_hard_delete_rate(self):
        self._as(30)
        rate_id = self.Session().query(Rate).one().id
        self.client.post("/economist/rates/delete", data={"rate_id": rate_id}, follow_redirects=False)
        self.assertIsNotNone(self.Session().get(Rate, rate_id))


RECHECK_HARNESS = r"""
const fs = require('fs');
const noop = () => {};
function field(name) { return { name, disabled: false }; }
const fields = [field('contract_number'), field('fix:le:inn'), field('fix:region:uvm_person'), field('fix:user:inn')];
const appended = [];
const form = { elements: fields, submitted: false, appendChild(el) { appended.push(el); }, submit() { this.submitted = true; } };
const select = { form, listeners: {}, getAttribute: n => n === 'data-recheck' ? process.argv[1] : null,
  addEventListener(t, f) { this.listeners[t] = f; } };
const stub = () => ({ classList: { add: noop, remove: noop, toggle: noop, contains: () => false }, style: {},
  setAttribute: noop, appendChild: noop, addEventListener: noop, querySelectorAll: () => [], querySelector: () => null });
global.window = { matchMedia: () => ({ matches: true, addEventListener: noop }), addEventListener: noop };
global.sessionStorage = { getItem: noop, setItem: noop };
global.document = { readyState: 'complete', body: stub(), documentElement: stub(), addEventListener: noop,
  createElement: () => ({}), getElementById: () => null, querySelector: () => null,
  querySelectorAll: sel => sel === 'form[data-recheck-form] select[data-recheck]' ? [select] : [] };
global.IntersectionObserver = function () { return { observe: noop }; };
eval(fs.readFileSync(process.argv[2], 'utf8'));
select.listeners.change();
console.log(JSON.stringify({ submitted: form.submitted, action: appended.map(e => [e.name, e.value]),
  disabled: fields.filter(f => f.disabled).map(f => f.name) }));
"""


@unittest.skipUnless(shutil.which("node"), "нужен node")
class RecheckJsTests(unittest.TestCase):
    def _run(self, kind):
        result = subprocess.run(["node", "-e", RECHECK_HARNESS, kind, UI_JS], capture_output=True, text=True, timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr)
        return json.loads(result.stdout)

    def test_entity_change_submits_check_without_old_entity_fixes(self):
        self.assertEqual(self._run("le"), {"submitted": True, "action": [["action", "check"]], "disabled": ["fix:le:inn"]})
        self.assertEqual(self._run("region")["disabled"], ["fix:region:uvm_person"])


if __name__ == "__main__":
    unittest.main()
