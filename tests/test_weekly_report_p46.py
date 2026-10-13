"""П4.6: «Еженедельный отчёт» — рабочий экран ТМ: строки отчёта, индивидуальная ставка
с понедельника, закрепление «ПУСТО», запрос смены ТМ, закрытая неделя, частичный аванс
в табеле за период, xlsx, права и выбор недели."""

import datetime
import os
import tempfile
import unittest
from io import BytesIO
from unittest import mock

os.environ.setdefault("DATABASE_URL", "sqlite:///:memory:")
os.environ.setdefault("SECRET_KEY", "test-secret")

import openpyxl
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from database import Base
from models import (AdvanceSelection, AuditLog, HourlyBonus, PayrollRun, PayrollRunItem, Rate, Requisite, Service,
                    Shift, Store, TmChangeRequest, User, WeeklyTmCheck)
from payouts_helpers import PARTIAL_ADVANCE_SERVICE, fix_advance, rebuild_period_items
from service_catalog import build_service_resolver
from utils import pick_rate
from weekly_actions import (WeeklyError, decide_tm_request, request_tm_change, set_advance, set_individual_rate,
                            set_week_closed, take_employee, toggle_advances_submitted)
from weekly_report import (NO_RATE, NO_TM, TOTAL, WEEKLY_SHEET, build_weekly_xlsx, resolve_week, weekly_filename,
                           weekly_report)

MONDAY = datetime.date(2026, 10, 5)
UNIVERSAL = "Универсальная услуга для ММФ"
KITCHEN = "Квалифицированные услуги в столовой"  # ставки нет
CARD = "4111111111111111"


def _session():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine)()


def _shift(employee, day, hours, service=UNIVERSAL, store="Лента-669 (г. Балашиха)", request_type="Основные заказы"):
    return Shift(employee=employee, store=store, city="Москва", format="СМ", service=f"3ур_{service}", hours=hours,
                 shift_date=MONDAY + datetime.timedelta(days=day), request_type=request_type)


def _fixture():
    s = _session()
    s.add_all([
        User(id=10, employee_name="Пиунов Андрей", role="tm"),
        User(id=12, employee_name="Амин", role="tm"),
        User(id=20, employee_name="Экономист", role="economist"),
        User(id=11, employee_name="Борисов Борис", role="employee", tm_id=10),
        User(id=13, employee_name="АРСТАНБЕКОВА Дамира", role="employee"),
        Service(name=UNIVERSAL), Service(name=KITCHEN),
        Store(tk_number=669, city="Балашиха", rate_region="Московская область", format="СМ"),
        Requisite(user_id=11, employee_name="Борисов Борис", card_number=CARD, is_active=True, is_verified=True),
    ])
    s.flush()
    universal = s.query(Service).filter_by(name=UNIVERSAL).one()
    s.add_all([
        Rate(service=f"3ур_{UNIVERSAL}", service_id=universal.id, level=3, format="СМ", city="Московская область",
             hourly_rate=333),
        HourlyBonus(name="Акция", amount_per_hour=5, active_from=datetime.date(2026, 9, 1), is_active=True),
        # Борисов: две услуги — две строки; столовая без ставки; смена без плана — красный 0.
        _shift("Борисов Борис", 0, 11), _shift("Борисов Борис", 1, 10),
        _shift("Борисов Борис", 2, 10, service=KITCHEN),
        _shift("Борисов Борис", 3, 8, request_type="Смена без плана"),
        # Регистр: «АРСТАНБЕКОВА» не встаёт первой; у «аблеевой» ещё нет users — «ПУСТО».
        _shift("АРСТАНБЕКОВА Дамира", 4, 11), _shift("аблеева Глюза", 5, 12),
        _shift("Вне недели", 7, 11),
    ])
    s.commit()
    return s


class WeeklyRowsTests(unittest.TestCase):
    def setUp(self):
        self.s = _fixture()
        self.report = weekly_report(self.s, MONDAY)
        self.rows = self.report["rows"]

    def test_rows(self):
        self.assertEqual([(r["name"], r["tk"], r["service"]) for r in self.rows], [
            ("аблеева Глюза", 669, f"3ур_{UNIVERSAL}"),
            ("АРСТАНБЕКОВА Дамира", 669, f"3ур_{UNIVERSAL}"),
            ("Борисов Борис", 669, f"3ур_{KITCHEN}"),
            ("Борисов Борис", 669, f"3ур_{UNIVERSAL}"),
        ])
        borisov = self.rows[3]
        self.assertEqual(borisov["days"][:4], [11, 10, None, 0])  # смена без плана — 0
        self.assertEqual(borisov["no_plan_days"], [False, False, False, True, False, False, False])
        self.assertEqual((borisov["hours"], borisov["rate"], borisov["bonus"], borisov["fot"]),
                         (21, 333, 21 * 5, 21 * 333 + 21 * 5))
        self.assertEqual((borisov["city"], borisov["format"], borisov["tm"]), ("Московская область", "СМ", "Пиунов Андрей"))
        self.assertTrue(borisov["has_requisites"])
        kitchen = self.rows[2]
        self.assertEqual((kitchen["rate"], kitchen["bonus"], kitchen["fot"], kitchen["no_rate"]), (None, 0, 0, True))
        self.assertEqual((self.rows[0]["tm"], self.rows[0]["user_id"], self.rows[0]["has_requisites"]), (NO_TM, None, False))
        self.assertEqual(self.report["people"]["борисов борис"]["fot"], 21 * 338)
        self.assertEqual((self.report["no_rate"]["shifts"], self.report["no_rate"]["hours"]), (1, 10))

    def test_xlsx(self):
        sheet = openpyxl.load_workbook(BytesIO(build_weekly_xlsx(self.report)))[WEEKLY_SHEET]
        header = [c.value for c in sheet[1]]
        self.assertEqual(header, ["ФИО сотрудника Исполнителя", "Номер ТК", "Город", "Тип", "Услуга",
                                  "5", "6", "7", "8", "9", "10", "11", "ИЧ", "ТМ", "Ставка", "Акция", "ФОТ"])
        self.assertTrue(sheet["A1"].font.b)
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet["I5"].value, 0)
        self.assertEqual(sheet["I5"].fill.fgColor.rgb, "00FF0000")  # без плана — красный
        self.assertEqual(sheet["A2"].fill.fgColor.rgb, "00FFC7CE")  # без реквизитов — розовая строка
        self.assertIsNone(sheet["A5"].fill.fill_type)
        self.assertEqual((sheet["N2"].value, sheet["O4"].value), (NO_TM, NO_RATE))
        self.assertEqual((sheet["O5"].value, sheet["P5"].value, sheet["Q5"].value), (333, 105, 21 * 338))
        last = [c.value for c in sheet[sheet.max_row]]
        self.assertEqual((last[0], last[12], last[16]), (TOTAL, 11 + 12 + 10 + 21, self.report["total_fot"]))
        self.assertEqual(weekly_filename(MONDAY), "Недельный_отчёт_05.10.2026-11.10.2026.xlsx")

    def test_resolve_week(self):
        today = datetime.date(2026, 10, 14)
        self.assertEqual(resolve_week("", today), MONDAY)
        self.assertEqual(resolve_week("2026-10-08", today), MONDAY)  # вниз до понедельника
        self.assertEqual(resolve_week("2026-09-28", today), datetime.date(2026, 9, 28))
        self.assertEqual(resolve_week("2026-10-12", today), MONDAY)  # незавершённая — недоступна
        self.assertEqual(resolve_week("2026-10-12", datetime.date(2026, 10, 19)), datetime.date(2026, 10, 12))


class WeeklyActionsTests(unittest.TestCase):
    def setUp(self):
        self.s = _fixture()
        self.tm = self.s.get(User, 10)
        self.tm2 = self.s.get(User, 12)
        self.economist = self.s.get(User, 20)

    def _pick(self, day):
        shift = _shift("Борисов Борис", day, 11)
        rate = pick_rate(self.s.query(Rate).all(), shift, resolve=build_service_resolver(self.s))
        return rate.hourly_rate if rate else None

    def test_individual_rate_from_monday_closes_previous(self):
        self.s.add(Rate(service=f"3ур_{UNIVERSAL}", employee_name="Борисов Борис", hourly_rate=340,
                        active_from=datetime.date(2026, 9, 1)))
        self.s.commit()
        set_individual_rate(self.s, None, self.tm, MONDAY, "Борисов Борис", f"3ур_{UNIVERSAL}", "350")
        self.s.commit()
        old = self.s.query(Rate).filter_by(hourly_rate=340).one()
        self.assertEqual(old.active_to, datetime.date(2026, 10, 4))
        new = self.s.query(Rate).filter_by(hourly_rate=350).one()
        self.assertEqual((new.active_from, new.employee_name, new.level), (MONDAY, "Борисов Борис", 3))
        self.assertIn("ТМ Пиунов Андрей", new.comment)
        self.assertEqual((self._pick(-1), self._pick(0), self._pick(14)), (340, 350, 350))
        self.assertTrue(weekly_report(self.s, MONDAY)["rows"][3]["rate_individual"])
        set_individual_rate(self.s, None, self.tm, MONDAY, "Борисов Борис", f"3ур_{UNIVERSAL}", "")  # к сетке
        self.s.commit()
        self.assertEqual(self._pick(0), 333)
        self.assertEqual(self.s.query(AuditLog).filter_by(action="rate.individual.set").count(), 2)
        with self.assertRaises(WeeklyError):  # чужой сотрудник
            set_individual_rate(self.s, None, self.tm2, MONDAY, "Борисов Борис", f"3ур_{UNIVERSAL}", "400")

    def test_take_empty_creates_user(self):
        employee = take_employee(self.s, None, self.tm, MONDAY, "аблеева Глюза")
        self.s.commit()
        self.assertEqual((employee.tm_id, employee.role, employee.password_hash), (10, "employee", None))
        self.assertEqual(weekly_report(self.s, MONDAY)["rows"][0]["tm"], "Пиунов Андрей")
        self.assertEqual(self.s.query(AuditLog).filter_by(action="user.tm.assign").count(), 1)

    def test_tm_change_request_and_accept(self):
        change, created = request_tm_change(self.s, None, self.tm2, MONDAY, "Борисов Борис")
        self.assertTrue(created)
        again, created = request_tm_change(self.s, None, self.tm2, MONDAY, "Борисов Борис")
        self.assertEqual((again.id, created), (change.id, False))
        with self.assertRaises(WeeklyError):  # ТМ запросы не решает
            decide_tm_request(self.s, None, self.tm, change.id, True)
        decide_tm_request(self.s, None, self.economist, change.id, True)
        self.s.commit()
        self.assertEqual(self.s.get(User, 11).tm_id, 12)
        self.assertEqual(self.s.query(TmChangeRequest).one().status, "accepted")

    def test_closed_week_blocks_tm_not_economist(self):
        set_week_closed(self.s, None, self.economist, MONDAY, True)
        self.s.commit()
        for action in (lambda: set_individual_rate(self.s, None, self.tm, MONDAY, "Борисов Борис",
                                                   f"3ур_{UNIVERSAL}", "350"),
                       lambda: take_employee(self.s, None, self.tm, MONDAY, "аблеева Глюза"),
                       lambda: toggle_advances_submitted(self.s, None, self.tm, MONDAY)):
            with self.assertRaises(WeeklyError):
                action()
        set_individual_rate(self.s, None, self.economist, MONDAY, "Борисов Борис", f"3ур_{UNIVERSAL}", "350")
        self.s.commit()
        self.assertEqual(self._pick(0), 350)

    def test_advance_selection_rules(self):
        people = weekly_report(self.s, MONDAY)["people"]
        with self.assertRaises(WeeklyError):  # нет реквизитов
            set_advance(self.s, None, self.economist, MONDAY, people["арстанбекова дамира"], True)
        with self.assertRaises(WeeklyError):  # больше ФОТ
            set_advance(self.s, None, self.tm, MONDAY, people["борисов борис"], True, "999999")
        selection = set_advance(self.s, None, self.tm, MONDAY, people["борисов борис"], True, "1000")
        self.assertEqual(selection.amount, 1000)
        set_advance(self.s, None, self.tm, MONDAY, people["борисов борис"], True)  # весь ФОТ
        self.assertIsNone(self.s.query(AdvanceSelection).one().amount)
        toggle_advances_submitted(self.s, None, self.tm, MONDAY)  # подал — заморожено для ТМ
        self.assertIsNotNone(self.s.query(WeeklyTmCheck).one().advances_submitted_at)
        with self.assertRaises(WeeklyError):
            set_advance(self.s, None, self.tm, MONDAY, people["борисов борис"], False)


class PartialAdvanceTests(unittest.TestCase):
    def test_partial_advance_deducted_in_period(self):
        s = _session()
        s.add_all([User(id=1, employee_name="Экономист", role="economist"),
                   User(id=11, employee_name="Борисов Борис", role="employee"),
                   Rate(service=f"3ур_{UNIVERSAL}", hourly_rate=1000),
                   _shift("Борисов Борис", 0, 10), _shift("Борисов Борис", 1, 10),
                   AdvanceSelection(week_start=MONDAY, employee_name="Борисов Борис", amount=10000)])
        s.commit()
        advance, errors = fix_advance(s, None, s.get(User, 1), MONDAY)
        self.assertEqual(errors, [])
        advance_items = s.query(PayrollRunItem).filter_by(run_id=advance.id).all()
        self.assertEqual([(i.service, i.shift_date, i.total_amount) for i in advance_items],
                         [(PARTIAL_ADVANCE_SERVICE, None, 10000)])
        period = PayrollRun(kind="period", status="draft", date_from=datetime.date(2026, 10, 1),
                            date_to=datetime.date(2026, 10, 15))
        s.add(period)
        s.flush()
        items = rebuild_period_items(s, period)
        shifts = [i for i in items if i.shift_date]
        self.assertEqual([(i.amount, i.paid_by_run_id, i.total_amount) for i in shifts],
                         [(10000, None, 10000), (10000, None, 10000)])  # строки смен полные
        partial = [i for i in items if not i.shift_date]
        self.assertEqual([(i.advanced_amount, i.total_amount, i.comment) for i in partial],
                         [(10000, -10000, "аванс 05.10–11.10 (частично)")])
        self.assertEqual(sum(i.total_amount for i in items), 10000)


class WeeklyRouteTests(unittest.TestCase):
    def setUp(self):
        import main
        from fastapi.testclient import TestClient

        from dependencies import current_user, get_db

        self.s = _fixture()
        self.app = main.app
        self.client = TestClient(main.app)
        self.app.dependency_overrides[get_db] = lambda: self.s
        self._current_user = current_user
        self.today = mock.patch("routers.weekly.business_today", return_value=datetime.date(2026, 10, 14))
        self.today.start()

    def tearDown(self):
        self.today.stop()
        self.app.dependency_overrides.clear()

    def _as(self, user_id):
        user = self.s.get(User, user_id)
        self.app.dependency_overrides[self._current_user] = lambda: user

    def test_tm_sees_own_by_default(self):
        self._as(10)
        html = self.client.get("/admin/payouts/weekly").text
        self.assertIn("05.10.2026 — 11.10.2026", html)
        self.assertIn("Борисов Борис", html)
        self.assertNotIn("аблеева Глюза", html)
        self.assertNotIn("Скачать xlsx", html)
        self.assertIn("аблеева Глюза", self.client.get("/admin/payouts/weekly?tm=0").text)  # все видят всех
        self.assertEqual(self.client.get("/admin/payouts/weekly?tm=").status_code, 200)  # пусто — по умолчанию «мои»
        response = self.client.post("/admin/payouts/weekly/rate", data={
            "week": "2026-10-05", "employee_name": "аблеева Глюза", "service": f"3ур_{UNIVERSAL}", "value": "400"})
        self.assertEqual(response.status_code, 403)  # не свой
        response = self.client.post("/admin/payouts/weekly/tm", data={
            "week": "2026-10-05", "employee_name": "аблеева Глюза", "action": "take"})
        self.assertEqual(response.json()["ok"], True)
        self.client.post("/admin/payouts/weekly/check", data={"week": "2026-10-05"})
        self.assertIsNotNone(self.s.query(WeeklyTmCheck).one().checked_at)

    def test_economist_export_and_close(self):
        self._as(20)
        html = self.client.get("/admin/payouts/weekly").text
        self.assertIn("аблеева Глюза", html)
        self.assertIn("Скачать xlsx", html)
        self.assertIn("Закрыто для коррекций", html)
        with tempfile.TemporaryDirectory() as tmp, mock.patch.dict(os.environ, {"GENERATED_FILES_ROOT": tmp}):
            response = self.client.get("/admin/payouts/weekly.xlsx?week=2026-10-07")
        self.assertEqual(openpyxl.load_workbook(BytesIO(response.content)).active.title, WEEKLY_SHEET)
        self.client.post("/admin/payouts/weekly/close", data={"week": "2026-10-05", "closed": "1"})
        self._as(10)
        response = self.client.post("/admin/payouts/weekly/rate", data={
            "week": "2026-10-05", "employee_name": "Борисов Борис", "service": f"3ур_{UNIVERSAL}", "value": "400"})
        self.assertEqual(response.status_code, 403)

    def test_dashboard_notices(self):
        request_tm_change(self.s, None, self.s.get(User, 12), MONDAY, "Борисов Борис")
        self.s.commit()
        self._as(20)
        with mock.patch("time_helpers.business_today", return_value=datetime.date(2026, 10, 14)):
            self.assertIn("Запросы на смену ТМ: 1", self.client.get("/economist").text)
            self._as(10)
            self.assertIn("Без ТМ в прошлой неделе: 2", self.client.get("/hr").text)

    def test_employee_has_no_access(self):
        self._as(11)
        response = self.client.get("/admin/payouts/weekly", follow_redirects=False)
        self.assertEqual((response.status_code, response.headers["location"]), (302, "/"))

    def test_page_shows_non_json_answer_instead_of_no_connection(self):
        # Прод: 500 text/plain на сохранении ставки показывался как «Нет связи». Теперь — статус и
        # начало тела; «Нет связи» — только когда fetch не дошёл до сервера.
        self._as(10)
        html = self.client.get("/admin/payouts/weekly").text
        self.assertIn("сервер ответил не JSON (HTTP ' + r.status", html)
        self.assertIn(".slice(0, 200)", html)
        self.assertIn("}, function () { alert('Нет связи — не сохранено'); return false; });", html)


class RateFormatNullableMigrationTests(unittest.TestCase):
    """Прод: rates.format в старой базе NOT NULL, модель — nullable; индивидуальная ставка недельного
    отчёта (без формата = на все форматы) падала NotNullViolation → 500 «Нет связи»."""

    def test_migration_drops_not_null_on_rates_format(self):
        import service_schema

        connection = mock.MagicMock()
        engine = mock.MagicMock()
        engine.begin.return_value.__enter__.return_value = connection
        with mock.patch.object(service_schema, "engine", engine):
            service_schema.ensure_service_columns()
        executed = [str(call.args[0]) for call in connection.execute.call_args_list]
        self.assertIn("ALTER TABLE rates ALTER COLUMN format DROP NOT NULL", executed)

    def test_model_keeps_format_nullable(self):
        self.assertTrue(Rate.__table__.c.format.nullable)


if __name__ == "__main__":
    unittest.main()
